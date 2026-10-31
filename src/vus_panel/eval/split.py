"""Deterministic calibration/test split of the Findlay BRCA1 variant IDs.

The answer key is the Findlay et al. 2018 saturation genome editing table,
read from data/raw/. It is never downloaded. Positions in that table are
hg19, so every variant here is GRCh37.

Each variant is placed by sha256(f"{seed}:{variant_id}") order within its
function class, so the split is stable across Python and numpy versions and
does not depend on input order.

Usage: python -m vus_panel.eval.split [--findlay PATH] [--out DIR] [--seed 0] [--calibration-fraction 0.2]
"""

from __future__ import annotations

import argparse
import hashlib
import json
import sys
from collections import defaultdict
from collections.abc import Iterable
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from vus_panel import config
from vus_panel.schemas import Variant

FINDLAY_FILENAME = "41586_2018_461_MOESM3_ESM.xlsx"
FINDLAY_SOURCE = "Findlay et al. 2018, Nature, 'Accurate classification of BRCA1 variants with saturation genome editing'"
HEADER_ROW = 3  # 1-based. Matches pandas header=2 in the Evo 2 BRCA1 notebook.
FINDLAY_BUILD = "GRCh37"  # the table's positions are hg19
COLUMNS = {
    "chrom": "chromosome",
    "pos": "position (hg19)",
    "ref": "reference",
    "alt": "alt",
    "score": "function.score.mean",
    "func_class": "func.class",
}
FUNCTION_CLASSES = ("FUNC", "INT", "LOF")
DEFAULT_SEED = 0
DEFAULT_CALIBRATION_FRACTION = 0.2
SPLIT_FILES = ("calibration.txt", "test.txt", "split_meta.json")


class FindlayFileMissing(FileNotFoundError):
    """The Findlay table is not in data/raw/."""


class FindlayFormatError(ValueError):
    """The Findlay table does not have the expected layout or contents."""


def default_findlay_path() -> Path:
    return config.raw_dir() / FINDLAY_FILENAME


def _text(value: Any) -> str | None:
    if value is None:
        return None
    return str(value).strip()


def _position(value: Any) -> int:
    if isinstance(value, bool):
        raise ValueError("position is not a number")
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, str) and value.strip().isdigit():
        return int(value.strip())
    raise ValueError(f"position {value!r} is not a whole number")


def _score(value: Any) -> float | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    return float(value)


def load_findlay(path: Path | None = None) -> list[Variant]:
    """Read the Findlay table. Header is the third row; columns are found by name."""
    path = Path(path) if path is not None else default_findlay_path()
    if not path.is_file():
        raise FindlayFileMissing(
            f"Findlay file not found at {config.display_path(path)}. Put {FINDLAY_FILENAME} "
            f"(supplementary data from {FINDLAY_SOURCE}) in data/raw/. This code never downloads it."
        )

    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = workbook.worksheets[0].iter_rows(values_only=True)
        header = None
        for row_number, row in enumerate(rows, start=1):
            if row_number == HEADER_ROW:
                header = row
                break
        if header is None:
            raise FindlayFormatError(f"{path.name} has fewer than {HEADER_ROW} rows")
        names = [_text(cell) or "" for cell in header]
        missing = [column for column in COLUMNS.values() if column not in names]
        if missing:
            raise FindlayFormatError(f"{path.name} header row {HEADER_ROW} is missing columns: {missing}")
        index = {key: names.index(column) for key, column in COLUMNS.items()}

        variants: list[Variant] = []
        seen: set[str] = set()
        for row_number, row in enumerate(rows, start=HEADER_ROW + 1):
            if all(cell is None or (isinstance(cell, str) and not cell.strip()) for cell in row):
                continue

            def cell(key: str) -> Any:
                i = index[key]
                return row[i] if i < len(row) else None

            func_class = _text(cell("func_class"))
            if func_class not in FUNCTION_CLASSES:
                raise FindlayFormatError(f"row {row_number}: unknown func.class {func_class!r}")
            try:
                variant = Variant(
                    chrom=_text(cell("chrom")),
                    pos=_position(cell("pos")),
                    ref=_text(cell("ref")),
                    alt=_text(cell("alt")),
                    genome_build=FINDLAY_BUILD,
                    findlay_function_class=func_class,
                    findlay_score=_score(cell("score")),
                )
            except (ValueError, TypeError) as exc:
                raise FindlayFormatError(f"row {row_number}: {exc}") from exc
            if variant.variant_id in seen:
                raise FindlayFormatError(f"row {row_number}: duplicate variant {variant.variant_id}")
            seen.add(variant.variant_id)
            variants.append(variant)
    finally:
        workbook.close()

    if not variants:
        raise FindlayFormatError(f"{path.name} has no variant rows")
    return variants


@dataclass(frozen=True)
class Split:
    calibration: list[str]
    test: list[str]
    counts: dict[str, dict[str, int]]
    seed: int
    calibration_fraction: float


def _order_key(seed: int, variant_id: str) -> str:
    return hashlib.sha256(f"{seed}:{variant_id}".encode("utf-8")).hexdigest()


def stratified_split(
    items: Iterable[tuple[str, str]],
    seed: int = DEFAULT_SEED,
    calibration_fraction: float = DEFAULT_CALIBRATION_FRACTION,
) -> Split:
    """Split (variant_id, function_class) pairs into calibration and test, stratified by class."""
    if not 0.0 < calibration_fraction < 1.0:
        raise ValueError("calibration_fraction must be strictly between 0 and 1")
    by_class: dict[str, list[str]] = defaultdict(list)
    seen: set[str] = set()
    for variant_id, function_class in items:
        if function_class not in FUNCTION_CLASSES:
            raise ValueError(f"{variant_id}: unknown function class {function_class!r}")
        if variant_id in seen:
            raise ValueError(f"duplicate variant ID {variant_id}")
        seen.add(variant_id)
        by_class[function_class].append(variant_id)
    if not seen:
        raise ValueError("nothing to split")

    calibration: list[str] = []
    test: list[str] = []
    counts: dict[str, dict[str, int]] = {}
    for function_class in sorted(by_class):
        ordered = sorted(by_class[function_class], key=lambda vid: (_order_key(seed, vid), vid))
        n_calibration = int(len(ordered) * calibration_fraction + 0.5)
        calibration.extend(ordered[:n_calibration])
        test.extend(ordered[n_calibration:])
        counts[function_class] = {"calibration": n_calibration, "test": len(ordered) - n_calibration}
    return Split(sorted(calibration), sorted(test), counts, seed, calibration_fraction)


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def write_split(split: Split, out_dir: Path, source_path: Path) -> None:
    """Write variant-ID lists and metadata. No scores or classes go into these files."""
    out_dir.mkdir(parents=True, exist_ok=True)
    (out_dir / "calibration.txt").write_text("".join(f"{v}\n" for v in split.calibration), encoding="utf-8")
    (out_dir / "test.txt").write_text("".join(f"{v}\n" for v in split.test), encoding="utf-8")
    meta = {
        "seed": split.seed,
        "calibration_fraction": split.calibration_fraction,
        "genome_build": FINDLAY_BUILD,
        "counts": split.counts,
        "n_calibration": len(split.calibration),
        "n_test": len(split.test),
        "source_file": source_path.name,
        "source_sha256": sha256_file(source_path),
    }
    (out_dir / "split_meta.json").write_text(json.dumps(meta, indent=2) + "\n", encoding="utf-8")


def read_ids(path: Path) -> set[str]:
    """Read a split file. One variant ID per line."""
    if not path.is_file():
        raise FileNotFoundError(
            f"split file not found at {config.display_path(path)}. Run python -m vus_panel.eval.split first."
        )
    return {line.strip() for line in path.read_text(encoding="utf-8").splitlines() if line.strip()}


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Split Findlay BRCA1 variant IDs into calibration and test sets.")
    parser.add_argument("--findlay", type=Path, default=None, help=f"default: data/raw/{FINDLAY_FILENAME}")
    parser.add_argument("--out", type=Path, default=None, help="default: data/splits/")
    parser.add_argument("--seed", type=int, default=DEFAULT_SEED)
    parser.add_argument("--calibration-fraction", type=float, default=DEFAULT_CALIBRATION_FRACTION)
    parser.add_argument("--force", action="store_true", help="overwrite existing split files")
    args = parser.parse_args(argv)

    source = args.findlay if args.findlay is not None else default_findlay_path()
    out_dir = args.out if args.out is not None else config.splits_dir()
    existing = [name for name in SPLIT_FILES if (out_dir / name).exists()]
    if existing and not args.force:
        print(
            f"error: {', '.join(existing)} already exist in {config.display_path(out_dir)}. Pass --force to overwrite.",
            file=sys.stderr,
        )
        return 2
    try:
        variants = load_findlay(source)
        split = stratified_split(
            ((v.variant_id, v.findlay_function_class) for v in variants),
            seed=args.seed,
            calibration_fraction=args.calibration_fraction,
        )
    except (FindlayFileMissing, ValueError) as exc:  # ValueError covers FindlayFormatError
        print(f"error: {exc}", file=sys.stderr)
        return 2
    write_split(split, out_dir, source)
    print(f"wrote {len(split.calibration)} calibration and {len(split.test)} test IDs to {config.display_path(out_dir)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
